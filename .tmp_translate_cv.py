from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

root = Path(__file__).parent
doc = Document(root / 'CV-Maynor-Lopez-Webmaster-Junior.docx')
texts = {
1: 'JUNIOR WEB DEVELOPER | FRONTEND DEVELOPMENT',
3: 'PROFESSIONAL SUMMARY',
4: 'Information Technology Engineering student building responsive web interfaces with HTML, CSS, JavaScript and React. Hands-on projects include interactive layouts, forms and REST API integration. Brings customer service, recruitment and project coordination experience, with strong communication and problem-solving skills. Eager to contribute to user-friendly websites, mobile compatibility and ongoing improvements.',
5: 'TECHNICAL SKILLS',
6: 'Frontend: HTML5, CSS3, JavaScript, React, Tailwind CSS and responsive design.',
7: 'Web development: DOM manipulation, REST API integration, form validation and basic accessibility.',
8: 'Tools & design: Git, GitHub, Vite, Visual Studio Code, Figma and basic UI/UX principles.',
9: 'Productivity: Advanced Excel, Microsoft Office, Google Workspace and basic Power BI.',
10: 'Languages: Native Spanish and advanced English.',
11: 'Currently learning: C#, .NET, ASP.NET Core and API design.',
12: 'SELECTED WEB PROJECTS',
14: 'Responsive product catalog with quantity controls, shopping cart and order confirmation.',
16: 'Interactive dashboard with filters, active/inactive states and theme switching.',
18: 'Responsive form with input validation, clear error feedback and submission confirmation.',
19: 'PROFESSIONAL EXPERIENCE',
21: 'Handle customer inquiries, troubleshoot issues and communicate clear solutions.',
22: 'Apply active listening and clear communication to support a positive customer experience.',
24: 'Sourced and screened candidates and conducted phone interviews for the CareerBuilder campaign.',
25: 'Coordinated interview confirmations, emails and follow-ups with attention to detail.',
27: 'Coordinated community education, health, infrastructure and recreation programs with stakeholders.',
28: 'Followed up on activities and community needs, strengthening teamwork and organizational skills.',
29: 'EDUCATION',
30: 'Information Technology Engineering (in progress) | UNAN-León | 2023–Present',
31: "Bachelor’s Degree in Business Administration | UNAN-León | 2012–2016",
32: 'Technical Diploma in Accounting | INATEC | 2014–2016',
33: 'ADDITIONAL TRAINING — PLATZI',
34: 'React.js — February 2026',
35: 'Professional REST API Consumption with JavaScript — August 2025',
36: 'Asynchronous JavaScript — June 2025',
37: 'JavaScript: DOM Manipulation — March 2025',
38: 'JavaScript Fundamentals — March 2025',
39: 'HTML and CSS in Depth — February 2025',
}
for i, text in texts.items():
    p = doc.paragraphs[i]
    runs = p.runs
    if 6 <= i <= 11:
        label, rest = text.split(':', 1)
        runs[0].text = label + ':'
        runs[1].text = rest
        for r in runs[2:]: r.text = ''
    else:
        runs[0].text = text
        for r in runs[1:]: r.text = ''

for i, name in [(13, 'Product List with Cart'), (15, 'Browser Extension Manager'), (17, 'Contact Form')]:
    doc.paragraphs[i].runs[0].text = name
for i, role, company, dates in [
    (20, 'Customer Service Representative', 'D&R', 'Nov 2022–Present'),
    (23, 'Professional Recruiter', 'Accedo Technologies', 'Nov 2020–Apr 2022'),
    (26, 'Social Responsibility Supervisor', 'Ingenio San Antonio', 'Jun 2017–Oct 2022'),
]:
    p = doc.paragraphs[i]
    p.runs[0].text = role
    p.runs[1].text = ' | ' + company
    p.runs[2].text = ' | ' + dates
    for r in p.runs[3:]: r.text = ''

for t in doc.element.xpath('//w:hyperlink//w:t'):
    if t.text == 'Portafolio': t.text = 'Portfolio'
    if t.text == 'Código': t.text = 'Code'
for lang in doc.element.xpath('//w:lang'):
    lang.set(qn('w:val'), 'en-US')
for p in doc.paragraphs:
    for r in p.runs:
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'en-US')
        r._r.get_or_add_rPr().append(lang)
doc.core_properties.title = 'Maynor López | Junior Web Developer'
doc.core_properties.subject = 'English CV tailored to the Web Developer position'
doc.core_properties.language = 'en-US'
dest = root / 'CV-Maynor-Lopez-Web-Developer-English.docx'
doc.save(dest)
print(dest)
